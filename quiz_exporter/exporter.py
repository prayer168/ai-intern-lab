import csv

import unicodedata

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HEADERS = ["題號", "題目", "選項A", "選項B", "選項C", "選項D", "答案"]
QUESTION_MIN_WIDTH = 40


def _display_width(value):
    """中文等全形字算 2 個字元寬，其他算 1 個。"""
    text = "" if value is None else str(value)
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def _style_sheet(ws):
    """標題列加粗、凍結第一列、依內容調整欄寬（題目欄至少 40）。"""
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col in range(1, ws.max_column + 1):
        width = max(_display_width(c.value) for c in ws[get_column_letter(col)]) + 2
        if col == 2:
            width = max(width, QUESTION_MIN_WIDTH)
        ws.column_dimensions[get_column_letter(col)].width = width


def _clean(value):
    """去掉 Excel 不允許的隱形控制字元（常見於從 Word / PDF 複製的文字），保留換行與 tab。"""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def export_to_excel(questions, out_path):
    """把題目清單匯出成 Excel，第一列是標題。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "題庫"
    ws.append(HEADERS)

    for i, q in enumerate(questions, start=1):
        options = list(q.get("options", [])) + [""] * 4
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=_clean(q["question"]))
        for col, opt in enumerate(options[:4], start=3):
            ws.cell(row=row, column=col, value=_clean(opt))
        ans = q.get("answer", "")
        ws.cell(row=row, column=7, value=ans.upper() if isinstance(ans, str) else ans)

    _style_sheet(ws)
    wb.save(out_path)
    return out_path


def export_to_csv(questions, out_path):
    """把題目清單匯出成 CSV（utf-8-sig，Windows Excel 開啟中文不亂碼），第一列是標題。"""
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS)
        for i, q in enumerate(questions, start=1):
            options = (list(q.get("options", [])) + [""] * 4)[:4]
            ans = q.get("answer", "")
            writer.writerow([i, q["question"], *options, ans.upper() if isinstance(ans, str) else ans])
    return out_path
