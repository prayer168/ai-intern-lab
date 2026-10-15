from openpyxl import load_workbook

from quiz_exporter import export_to_excel, load_questions


def test_export_creates_file_with_headers(tmp_path):
    out = tmp_path / "out.xlsx"
    export_to_excel(load_questions("data/sample_quiz.json"), out)
    ws = load_workbook(out).active
    assert ws.title == "題庫"
    assert ws.cell(row=1, column=1).value == "題號"
    assert ws.cell(row=1, column=7).value == "答案"


def test_export_contains_all_questions(tmp_path):
    out = tmp_path / "out.xlsx"
    questions = load_questions("data/sample_quiz.json")
    export_to_excel(questions, out)
    ws = load_workbook(out).active
    texts = [c.value for c in ws["B"] if c.value]
    for q in questions:
        assert q["question"] in texts


def test_export_first_question_on_row_2(tmp_path):
    out = tmp_path / "out.xlsx"
    questions = load_questions("data/sample_quiz.json")
    export_to_excel(questions, out)
    ws = load_workbook(out).active
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == questions[0]["question"]


def test_export_lowercase_answer_unified_to_uppercase(tmp_path):
    out = tmp_path / "out.xlsx"
    questions = [{"question": "題目一", "options": ["A", "B", "C", "D"], "answer": "c"}]
    export_to_excel(questions, out)
    ws = load_workbook(out).active
    assert ws.cell(row=2, column=7).value == "C"


def test_export_fewer_options_pads_empty_cells(tmp_path):
    import json

    p = tmp_path / "q.json"
    p.write_text(
        json.dumps({"questions": [{"question": "Short", "options": ["One", "Two"], "answer": "A"}]}),
        encoding="utf-8",
    )
    out = tmp_path / "out.xlsx"
    export_to_excel(load_questions(p), out)
    ws = load_workbook(out).active
    assert [ws.cell(row=2, column=c).value for c in range(3, 7)] == ["One", "Two", None, None]


def test_export_header_row_is_bold(tmp_path):
    out = tmp_path / "out.xlsx"
    export_to_excel(load_questions("data/sample_quiz.json"), out)
    ws = load_workbook(out).active
    assert all(ws.cell(row=1, column=c).font.bold for c in range(1, 8))
    assert not ws.cell(row=2, column=2).font.bold


def test_export_freezes_header_row(tmp_path):
    out = tmp_path / "out.xlsx"
    export_to_excel(load_questions("data/sample_quiz.json"), out)
    ws = load_workbook(out).active
    assert ws.freeze_panes == "A2"


def test_export_question_column_at_least_40_wide(tmp_path):
    out = tmp_path / "out.xlsx"
    export_to_excel([{"question": "短", "options": ["A"], "answer": "A"}], out)
    ws = load_workbook(out).active
    assert ws.column_dimensions["B"].width >= 40


def test_export_other_columns_fit_content_with_cjk_as_double_width(tmp_path):
    out = tmp_path / "out.xlsx"
    export_to_excel([{"question": "題", "options": ["二氧化碳二氧化碳", "B"], "answer": "A"}], out)
    ws = load_workbook(out).active
    # 「二氧化碳二氧化碳」8 個中文字 = 16 個字元寬
    assert ws.column_dimensions["C"].width >= 16
    # 「選項B」標題 = 2*2 + 1 = 5 個字元寬
    assert ws.column_dimensions["D"].width >= 5


def test_invisible_control_characters_are_removed(tmp_path):
    out = tmp_path / "out.xlsx"
    # 從 Word / PDF 複製過來的文字常夾帶 \x0b（垂直定位字元）等 Excel 不允許的控制字元
    export_to_excel([{"question": "水在\x0b幾度沸騰？", "options": ["100\x01°C", "第一行\n第二行"], "answer": "A"}], out)
    ws = load_workbook(out).active
    assert ws.cell(row=2, column=2).value == "水在幾度沸騰？"
    assert ws.cell(row=2, column=3).value == "100°C"
    assert ws.cell(row=2, column=4).value == "第一行\n第二行"
